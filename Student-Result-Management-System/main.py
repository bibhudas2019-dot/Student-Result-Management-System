import tkinter as tk
from tkinter import ttk, messagebox
from openpyxl import Workbook, load_workbook
import os


# ============================================================
# CONFIGURATION
# ============================================================

FILE_NAME = "student_results.xlsx"
PASS_MARK = 35
SUBJECTS = ["Python", "Computer Networks", "DBMS", "Mathematics", "Web Technology"]


# ============================================================
# EXCEL FUNCTIONS
# ============================================================

def create_excel_file():
    """Create the Excel file if it does not already exist."""

    if not os.path.exists(FILE_NAME):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Student Results"

        sheet.append([
            "Name",
            "Roll No",
            "Class",
            SUBJECTS[0],
            SUBJECTS[1],
            SUBJECTS[2],
            SUBJECTS[3],
            SUBJECTS[4],
            "Total",
            "Percentage",
            "Result"
        ])

        workbook.save(FILE_NAME)


def read_all_students():
    """Read all student records from Excel."""

    create_excel_file()

    workbook = load_workbook(FILE_NAME)
    sheet = workbook.active

    students = []

    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row[0] is not None:
            students.append(row)

    workbook.close()

    return students


# ============================================================
# UTILITY FUNCTIONS
# ============================================================

def clear_add_form():
    name_entry.delete(0, tk.END)
    roll_entry.delete(0, tk.END)
    class_entry.delete(0, tk.END)

    for entry in mark_entries:
        entry.delete(0, tk.END)

    name_entry.focus()


def update_dashboard():
    """Update total, passed and failed student counts."""

    students = read_all_students()

    total_students = len(students)
    passed_students = 0
    failed_students = 0

    for student in students:
        if str(student[10]) == "Pass":
            passed_students += 1
        else:
            failed_students += 1

    total_count.config(text=str(total_students))
    pass_count.config(text=str(passed_students))
    fail_count.config(text=str(failed_students))


def show_page(page):
    """Show one page and hide the others."""

    dashboard_page.pack_forget()
    add_page.pack_forget()
    search_page.pack_forget()
    results_page.pack_forget()

    page.pack(fill="both", expand=True)

    if page == dashboard_page:
        update_dashboard()

    elif page == results_page:
        load_all_results()


# ============================================================
# ADD STUDENT
# ============================================================

def save_student():
    """Validate and save a student record."""

    name = name_entry.get().strip()
    roll = roll_entry.get().strip()
    student_class = class_entry.get().strip()

    # Check basic information
    if name == "" or roll == "" or student_class == "":
        messagebox.showerror(
            "Missing Information",
            "Please enter Name, Roll Number and Class."
        )
        return

    # Check roll number
    if not roll.isdigit():
        messagebox.showerror(
            "Invalid Roll Number",
            "Roll Number must contain numbers only."
        )
        return

    # Read marks
    marks = []

    for entry in mark_entries:
        value = entry.get().strip()

        if value == "":
            messagebox.showerror(
                "Missing Marks",
                "Please enter marks for all subjects."
            )
            return

        if not value.isdigit():
            messagebox.showerror(
                "Invalid Marks",
                "Marks must be numbers between 0 and 100."
            )
            return

        mark = int(value)

        if mark < 0 or mark > 100:
            messagebox.showerror(
                "Invalid Marks",
                "Marks must be between 0 and 100."
            )
            return

        marks.append(mark)

    # Check duplicate roll number
    students = read_all_students()

    for student in students:
        if str(student[1]) == roll:
            messagebox.showerror(
                "Duplicate Roll Number",
                "This Roll Number already exists."
            )
            return

    # Calculate result
    total = sum(marks)
    percentage = total / len(SUBJECTS)

    if all(mark >= PASS_MARK for mark in marks):
        result = "Pass"
    else:
        result = "Fail"

    # Save to Excel
    workbook = load_workbook(FILE_NAME)
    sheet = workbook.active

    sheet.append([
        name,
        int(roll),
        student_class,
        marks[0],
        marks[1],
        marks[2],
        marks[3],
        marks[4],
        total,
        round(percentage, 2),
        result
    ])

    workbook.save(FILE_NAME)
    workbook.close()

    messagebox.showinfo(
        "Success",
        "Student record saved successfully!"
    )

    clear_add_form()
    update_dashboard()


# ============================================================
# SEARCH RESULT
# ============================================================

def search_student():
    """Search a student using Roll Number."""

    roll = search_roll_entry.get().strip()

    if roll == "":
        messagebox.showerror(
            "Search Error",
            "Please enter a Roll Number."
        )
        return

    students = read_all_students()

    for student in students:

        if str(student[1]) == roll:

            search_name_value.config(text=str(student[0]))
            search_roll_value.config(text=str(student[1]))
            search_class_value.config(text=str(student[2]))
            search_total_value.config(text=str(student[8]))
            search_percentage_value.config(
                text=str(student[9]) + "%"
            )
            search_result_value.config(
                text=str(student[10]),
                foreground="#16803C" if student[10] == "Pass" else "#D62828"
            )

            return

    search_name_value.config(text="—")
    search_roll_value.config(text="—")
    search_class_value.config(text="—")
    search_total_value.config(text="—")
    search_percentage_value.config(text="—")
    search_result_value.config(
        text="Student Not Found",
        foreground="#D62828"
    )


def clear_search():
    search_roll_entry.delete(0, tk.END)

    search_name_value.config(text="—")
    search_roll_value.config(text="—")
    search_class_value.config(text="—")
    search_total_value.config(text="—")
    search_percentage_value.config(text="—")
    search_result_value.config(
        text="—",
        foreground="#202A44"
    )


# ============================================================
# SHOW ALL RESULTS
# ============================================================

def load_all_results():
    """Display every student in the Treeview."""

    for item in result_table.get_children():
        result_table.delete(item)

    students = read_all_students()

    for student in students:

        result_table.insert(
            "",
            tk.END,
            values=(
                student[0],
                student[1],
                student[2],
                student[8],
                str(student[9]) + "%",
                student[10]
            )
        )


# ============================================================
# MAIN WINDOW
# ============================================================

create_excel_file()

root = tk.Tk()
root.title("Student Result Management System")
root.geometry("1100x700")
root.minsize(950, 600)
root.configure(bg="#F4F7FB")


# ============================================================
# COLORS
# ============================================================

SIDEBAR_COLOR = "#202A44"
ACCENT_COLOR = "#4F46E5"
LIGHT_BG = "#F4F7FB"
CARD_COLOR = "#FFFFFF"
TEXT_COLOR = "#202A44"
MUTED_COLOR = "#6B7280"
SUCCESS_COLOR = "#16803C"
DANGER_COLOR = "#D62828"


# ============================================================
# STYLE
# ============================================================

style = ttk.Style()
style.theme_use("clam")

style.configure(
    "Treeview",
    background="white",
    foreground=TEXT_COLOR,
    rowheight=35,
    fieldbackground="white",
    font=("Segoe UI", 10)
)

style.configure(
    "Treeview.Heading",
    background=SIDEBAR_COLOR,
    foreground="white",
    font=("Segoe UI", 10, "bold"),
    padding=8
)

style.map(
    "Treeview",
    background=[("selected", "#E0E7FF")],
    foreground=[("selected", TEXT_COLOR)]
)


# ============================================================
# SIDEBAR
# ============================================================

sidebar = tk.Frame(
    root,
    bg=SIDEBAR_COLOR,
    width=230
)

sidebar.pack(
    side="left",
    fill="y"
)

sidebar.pack_propagate(False)


# Logo / Title

tk.Label(
    sidebar,
    text="SRMS",
    font=("Segoe UI", 28, "bold"),
    bg=SIDEBAR_COLOR,
    fg="white"
).pack(pady=(35, 2))

tk.Label(
    sidebar,
    text="Student Result\nManagement",
    font=("Segoe UI", 10),
    bg=SIDEBAR_COLOR,
    fg="#B8C2D9",
    justify="center"
).pack(pady=(0, 35))


# Sidebar buttons

def sidebar_button(text, command):
    button = tk.Button(
        sidebar,
        text=text,
        command=command,
        font=("Segoe UI", 11, "bold"),
        bg=SIDEBAR_COLOR,
        fg="white",
        activebackground=ACCENT_COLOR,
        activeforeground="white",
        relief="flat",
        bd=0,
        anchor="w",
        padx=25,
        pady=13,
        cursor="hand2"
    )

    button.pack(fill="x", padx=12, pady=3)

    return button


sidebar_button(
    "⌂   Dashboard",
    lambda: show_page(dashboard_page)
)

sidebar_button(
    "＋   Add Student",
    lambda: show_page(add_page)
)

sidebar_button(
    "⌕   Search Result",
    lambda: show_page(search_page)
)

sidebar_button(
    "▦   All Results",
    lambda: show_page(results_page)
)


# Bottom exit button

tk.Button(
    sidebar,
    text="Exit Application",
    command=root.destroy,
    font=("Segoe UI", 10, "bold"),
    bg="#2D3856",
    fg="#FFB4B4",
    activebackground="#D62828",
    activeforeground="white",
    relief="flat",
    cursor="hand2",
    pady=10
).pack(
    side="bottom",
    fill="x",
    padx=20,
    pady=25
)


# ============================================================
# CONTENT AREA
# ============================================================

content = tk.Frame(
    root,
    bg=LIGHT_BG
)

content.pack(
    side="right",
    fill="both",
    expand=True
)


# ============================================================
# DASHBOARD PAGE
# ============================================================

dashboard_page = tk.Frame(
    content,
    bg=LIGHT_BG
)

# Header

tk.Label(
    dashboard_page,
    text="Dashboard",
    font=("Segoe UI", 25, "bold"),
    bg=LIGHT_BG,
    fg=TEXT_COLOR
).pack(
    anchor="w",
    padx=35,
    pady=(30, 5)
)

tk.Label(
    dashboard_page,
    text="Monitor student performance at a glance.",
    font=("Segoe UI", 11),
    bg=LIGHT_BG,
    fg=MUTED_COLOR
).pack(
    anchor="w",
    padx=35
)


# Statistics cards

stats_frame = tk.Frame(
    dashboard_page,
    bg=LIGHT_BG
)

stats_frame.pack(
    fill="x",
    padx=35,
    pady=30
)


def create_stat_card(parent, title):
    card = tk.Frame(
        parent,
        bg=CARD_COLOR,
        height=140
    )

    card.pack(
        side="left",
        fill="x",
        expand=True,
        padx=8
    )

    tk.Label(
        card,
        text=title,
        font=("Segoe UI", 10, "bold"),
        bg=CARD_COLOR,
        fg=MUTED_COLOR
    ).pack(
        anchor="w",
        padx=22,
        pady=(22, 5)
    )

    value = tk.Label(
        card,
        text="0",
        font=("Segoe UI", 30, "bold"),
        bg=CARD_COLOR,
        fg=TEXT_COLOR
    )

    value.pack(
        anchor="w",
        padx=22
    )

    return value


total_count = create_stat_card(
    stats_frame,
    "TOTAL STUDENTS"
)

pass_count = create_stat_card(
    stats_frame,
    "PASSED STUDENTS"
)

fail_count = create_stat_card(
    stats_frame,
    "FAILED STUDENTS"
)


# Welcome section

welcome_card = tk.Frame(
    dashboard_page,
    bg=CARD_COLOR
)

welcome_card.pack(
    fill="both",
    expand=True,
    padx=43,
    pady=5
)

tk.Label(
    welcome_card,
    text="Welcome to SRMS",
    font=("Segoe UI", 22, "bold"),
    bg=CARD_COLOR,
    fg=TEXT_COLOR
).pack(
    anchor="w",
    padx=30,
    pady=(35, 10)
)

tk.Label(
    welcome_card,
    text=(
        "A simple digital solution for maintaining student academic records.\n\n"
        "You can add new students, calculate their performance automatically,\n"
        "search individual results and view all records in one place."
    ),
    font=("Segoe UI", 11),
    bg=CARD_COLOR,
    fg=MUTED_COLOR,
    justify="left"
).pack(
    anchor="w",
    padx=30
)

tk.Button(
    welcome_card,
    text="＋  Add Your First Student",
    command=lambda: show_page(add_page),
    font=("Segoe UI", 10, "bold"),
    bg=ACCENT_COLOR,
    fg="white",
    activebackground="#3730A3",
    activeforeground="white",
    relief="flat",
    cursor="hand2",
    padx=18,
    pady=10
).pack(
    anchor="w",
    padx=30,
    pady=25
)


# ============================================================
# ADD STUDENT PAGE
# ============================================================

add_page = tk.Frame(
    content,
    bg=LIGHT_BG
)

tk.Label(
    add_page,
    text="Add Student",
    font=("Segoe UI", 25, "bold"),
    bg=LIGHT_BG,
    fg=TEXT_COLOR
).pack(
    anchor="w",
    padx=35,
    pady=(30, 5)
)

tk.Label(
    add_page,
    text="Enter student information and subject marks.",
    font=("Segoe UI", 11),
    bg=LIGHT_BG,
    fg=MUTED_COLOR
).pack(
    anchor="w",
    padx=35
)


form_card = tk.Frame(
    add_page,
    bg=CARD_COLOR
)

form_card.pack(
    fill="both",
    expand=True,
    padx=35,
    pady=25
)


# Student information heading

tk.Label(
    form_card,
    text="Student Information",
    font=("Segoe UI", 14, "bold"),
    bg=CARD_COLOR,
    fg=TEXT_COLOR
).grid(
    row=0,
    column=0,
    columnspan=4,
    sticky="w",
    padx=30,
    pady=(25, 15)
)


# Form labels and entries

tk.Label(
    form_card,
    text="Student Name",
    font=("Segoe UI", 10, "bold"),
    bg=CARD_COLOR,
    fg=MUTED_COLOR
).grid(
    row=1,
    column=0,
    sticky="w",
    padx=30,
    pady=5
)

name_entry = tk.Entry(
    form_card,
    font=("Segoe UI", 11),
    relief="solid",
    bd=1
)

name_entry.grid(
    row=2,
    column=0,
    sticky="ew",
    padx=30,
    pady=(0, 15),
    ipady=7
)


tk.Label(
    form_card,
    text="Roll Number",
    font=("Segoe UI", 10, "bold"),
    bg=CARD_COLOR,
    fg=MUTED_COLOR
).grid(
    row=1,
    column=1,
    sticky="w",
    padx=15,
    pady=5
)

roll_entry = tk.Entry(
    form_card,
    font=("Segoe UI", 11),
    relief="solid",
    bd=1
)

roll_entry.grid(
    row=2,
    column=1,
    sticky="ew",
    padx=15,
    pady=(0, 15),
    ipady=7
)


tk.Label(
    form_card,
    text="Class",
    font=("Segoe UI", 10, "bold"),
    bg=CARD_COLOR,
    fg=MUTED_COLOR
).grid(
    row=1,
    column=2,
    sticky="w",
    padx=15,
    pady=5
)

class_entry = tk.Entry(
    form_card,
    font=("Segoe UI", 11),
    relief="solid",
    bd=1
)

class_entry.grid(
    row=2,
    column=2,
    sticky="ew",
    padx=15,
    pady=(0, 15),
    ipady=7
)


# Subject section

tk.Label(
    form_card,
    text="Subject Marks",
    font=("Segoe UI", 14, "bold"),
    bg=CARD_COLOR,
    fg=TEXT_COLOR
).grid(
    row=3,
    column=0,
    columnspan=4,
    sticky="w",
    padx=30,
    pady=(15, 10)
)


mark_entries = []

for i, subject in enumerate(SUBJECTS):

    row_number = 4 + (i // 3)
    column_number = i % 3

    subject_box = tk.Frame(
        form_card,
        bg=CARD_COLOR
    )

    subject_box.grid(
        row=row_number,
        column=column_number,
        sticky="ew",
        padx=30 if column_number == 0 else 15,
        pady=8
    )

    tk.Label(
        subject_box,
        text=subject,
        font=("Segoe UI", 10, "bold"),
        bg=CARD_COLOR,
        fg=MUTED_COLOR
    ).pack(
        anchor="w"
    )

    entry = tk.Entry(
        subject_box,
        font=("Segoe UI", 11),
        relief="solid",
        bd=1
    )

    entry.pack(
        fill="x",
        pady=5,
        ipady=7
    )

    mark_entries.append(entry)


# Make columns expand

for column in range(3):
    form_card.grid_columnconfigure(
        column,
        weight=1
    )


# Buttons

button_area = tk.Frame(
    form_card,
    bg=CARD_COLOR
)

button_area.grid(
    row=7,
    column=0,
    columnspan=4,
    sticky="w",
    padx=30,
    pady=25
)


tk.Button(
    button_area,
    text="SAVE STUDENT",
    command=save_student,
    font=("Segoe UI", 10, "bold"),
    bg=ACCENT_COLOR,
    fg="white",
    activebackground="#3730A3",
    activeforeground="white",
    relief="flat",
    cursor="hand2",
    padx=25,
    pady=10
).pack(
    side="left",
    padx=(0, 12)
)


tk.Button(
    button_area,
    text="CLEAR FORM",
    command=clear_add_form,
    font=("Segoe UI", 10, "bold"),
    bg="#E5E7EB",
    fg=TEXT_COLOR,
    activebackground="#D1D5DB",
    relief="flat",
    cursor="hand2",
    padx=25,
    pady=10
).pack(
    side="left"
)


# ============================================================
# SEARCH PAGE
# ============================================================

search_page = tk.Frame(
    content,
    bg=LIGHT_BG
)

tk.Label(
    search_page,
    text="Search Result",
    font=("Segoe UI", 25, "bold"),
    bg=LIGHT_BG,
    fg=TEXT_COLOR
).pack(
    anchor="w",
    padx=35,
    pady=(30, 5)
)

tk.Label(
    search_page,
    text="Find a student's result using their Roll Number.",
    font=("Segoe UI", 11),
    bg=LIGHT_BG,
    fg=MUTED_COLOR
).pack(
    anchor="w",
    padx=35
)


search_card = tk.Frame(
    search_page,
    bg=CARD_COLOR
)

search_card.pack(
    fill="both",
    expand=True,
    padx=35,
    pady=25
)


# Search input

tk.Label(
    search_card,
    text="Enter Roll Number",
    font=("Segoe UI", 11, "bold"),
    bg=CARD_COLOR,
    fg=TEXT_COLOR
).pack(
    anchor="w",
    padx=35,
    pady=(35, 8)
)


search_roll_entry = tk.Entry(
    search_card,
    font=("Segoe UI", 12),
    relief="solid",
    bd=1
)

search_roll_entry.pack(
    anchor="w",
    padx=35,
    ipadx=10,
    ipady=8
)


search_buttons = tk.Frame(
    search_card,
    bg=CARD_COLOR
)

search_buttons.pack(
    anchor="w",
    padx=35,
    pady=15
)


tk.Button(
    search_buttons,
    text="SEARCH RESULT",
    command=search_student,
    font=("Segoe UI", 10, "bold"),
    bg=ACCENT_COLOR,
    fg="white",
    activebackground="#3730A3",
    activeforeground="white",
    relief="flat",
    cursor="hand2",
    padx=20,
    pady=9
).pack(
    side="left",
    padx=(0, 10)
)


tk.Button(
    search_buttons,
    text="CLEAR",
    command=clear_search,
    font=("Segoe UI", 10, "bold"),
    bg="#E5E7EB",
    fg=TEXT_COLOR,
    activebackground="#D1D5DB",
    relief="flat",
    cursor="hand2",
    padx=20,
    pady=9
).pack(
    side="left"
)


# Result card

result_card = tk.Frame(
    search_card,
    bg="#F8FAFC",
    relief="solid",
    bd=1
)

result_card.pack(
    fill="x",
    padx=35,
    pady=20
)


tk.Label(
    result_card,
    text="RESULT DETAILS",
    font=("Segoe UI", 11, "bold"),
    bg="#F8FAFC",
    fg=MUTED_COLOR
).grid(
    row=0,
    column=0,
    columnspan=2,
    sticky="w",
    padx=25,
    pady=(20, 15)
)


def result_row(row, title):
    tk.Label(
        result_card,
        text=title,
        font=("Segoe UI", 10, "bold"),
        bg="#F8FAFC",
        fg=MUTED_COLOR
    ).grid(
        row=row,
        column=0,
        sticky="w",
        padx=25,
        pady=7
    )

    value = tk.Label(
        result_card,
        text="—",
        font=("Segoe UI", 10, "bold"),
        bg="#F8FAFC",
        fg=TEXT_COLOR
    )

    value.grid(
        row=row,
        column=1,
        sticky="w",
        padx=40,
        pady=7
    )

    return value


search_name_value = result_row(1, "Student Name")
search_roll_value = result_row(2, "Roll Number")
search_class_value = result_row(3, "Class")
search_total_value = result_row(4, "Total Marks")
search_percentage_value = result_row(5, "Percentage")
search_result_value = result_row(6, "Result")


# ============================================================
# ALL RESULTS PAGE
# ============================================================

results_page = tk.Frame(
    content,
    bg=LIGHT_BG
)

tk.Label(
    results_page,
    text="All Results",
    font=("Segoe UI", 25, "bold"),
    bg=LIGHT_BG,
    fg=TEXT_COLOR
).pack(
    anchor="w",
    padx=35,
    pady=(30, 5)
)

tk.Label(
    results_page,
    text="Complete student performance records.",
    font=("Segoe UI", 11),
    bg=LIGHT_BG,
    fg=MUTED_COLOR
).pack(
    anchor="w",
    padx=35
)


table_card = tk.Frame(
    results_page,
    bg=CARD_COLOR
)

table_card.pack(
    fill="both",
    expand=True,
    padx=35,
    pady=25
)


columns = (
    "Name",
    "Roll No",
    "Class",
    "Total Marks",
    "Percentage",
    "Result"
)


result_table = ttk.Treeview(
    table_card,
    columns=columns,
    show="headings"
)


for column in columns:

    result_table.heading(
        column,
        text=column
    )

    result_table.column(
        column,
        anchor="center",
        width=150
    )


result_table.pack(
    fill="both",
    expand=True,
    padx=20,
    pady=(20, 10)
)


# Scrollbar

scrollbar = ttk.Scrollbar(
    table_card,
    orient="vertical",
    command=result_table.yview
)

result_table.configure(
    yscrollcommand=scrollbar.set
)

scrollbar.pack(
    side="right",
    fill="y"
)


tk.Button(
    table_card,
    text="↻  REFRESH RESULTS",
    command=load_all_results,
    font=("Segoe UI", 10, "bold"),
    bg=ACCENT_COLOR,
    fg="white",
    activebackground="#3730A3",
    activeforeground="white",
    relief="flat",
    cursor="hand2",
    padx=20,
    pady=9
).pack(
    anchor="e",
    padx=20,
    pady=(0, 15)
)


# ============================================================
# START APPLICATION
# ============================================================

show_page(dashboard_page)
update_dashboard()

root.mainloop()